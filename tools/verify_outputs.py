"""Independent source recomputation; read-only SQLite, XLSX and PDF checks."""
import hashlib
import json
import math
import sqlite3
import zipfile
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parents[1]
a=json.loads((ROOT/'results/analysis.json').read_text())
m=a['metrics']
columns=['Complaint ID','Date received','Date sent to company','Submitted via','Timely response?']
d=pd.read_csv(ROOT/'data/raw/Consumer_Complaints.csv',usecols=columns,dtype=str,keep_default_na=False)
received=pd.to_datetime(d['Date received'],format='%m/%d/%Y',errors='coerce')
sent=pd.to_datetime(d['Date sent to company'],format='%m/%d/%Y',errors='coerce')
days=(sent-received).dt.days
valid=days[days>=0].sort_values()
yes=(d['Timely response?'].str.strip().str.lower()=='yes')
no=(d['Timely response?'].str.strip().str.lower()=='no')
flag=(days<0)|days.isna()|(days>m['review_days'])|~yes
hasher=hashlib.sha256()
with (ROOT/'data/raw/Consumer_Complaints.csv').open('rb') as stream:
    for chunk in iter(lambda:stream.read(1024*1024),b''):hasher.update(chunk)
checks={
 'source_fingerprint':hasher.hexdigest()==a['source']['sha256'],
 'source_rows':len(d)==m['input_rows'],
 'unique_source_ids':d['Complaint ID'].nunique()==m['cases'],
 'timely_yes':int(yes.sum())==m['timely_yes'],
 'timely_no':int(no.sum())==m['timely_no'],
 'negative_intervals':int((days<0).sum())==m['negative_intervals'],
 'valid_intervals':len(valid)==m['valid_intervals'],
 'median_days':float(valid.median())==m['median_intake_days'],
 'p90_days':int(valid.iloc[math.ceil(.9*len(valid))-1])==m['p90_intake_days'],
 'routing_over_threshold':int((days>m['review_days']).sum())==m['over_review_days'],
 'review_queue_count':int(flag.sum())==m['review_queue_cases'],
}
expected=d.groupby('Submitted via').size().to_dict()
checks['channel_counts']=expected=={x['channel']:x['cases'] for x in a['reports']['channels']}
with sqlite3.connect(ROOT/'results/complaints.sqlite') as db:
    checks['sqlite_integrity']=db.execute('PRAGMA integrity_check').fetchone()[0]=='ok'
    checks['source_lineage']=db.execute('SELECT COUNT(*) FROM cases c LEFT JOIN selected s USING(source_row) WHERE s.source_row IS NULL').fetchone()[0]==0
    checks['omitted_narratives']='Consumer complaint narrative' not in str(db.execute('PRAGMA table_info(selected)').fetchall())
book=ROOT/'deliverables/complaint-operations-review.xlsx'
values=load_workbook(book,data_only=True)
formulas=load_workbook(book,data_only=False)
checks['xlsx_count']=values['Operations review']['B6'].value==m['cases']
checks['xlsx_rate']=abs(values['Operations review']['B9'].value-m['timely_rate'])<1e-10
checks['xlsx_formula']=formulas['Operations review']['B9'].data_type=='f'
checks['xlsx_review_rows']=values['Review queue'].max_row>=105 and values['Review queue']['A105'].value is not None
checks['xlsx_notes_blank']=all(not values['Review queue'].cell(row,col).value for row in range(6,106) for col in (7,9))
checks['xlsx_resolution_guard']='AND(' in formulas['Review queue']['J6'].value
checks['xlsx_status_validation']=len(formulas['Review queue'].data_validations.dataValidation)==1
with zipfile.ZipFile(book) as z:
    chart=z.read(next(n for n in z.namelist() if n.endswith('/chart1.xml'))).decode()
    checks['xlsx_native_chart']='Monthly' in chart and '$L$6:$L$17' in chart
pdf=PdfReader(ROOT/'deliverables/case-study.pdf',strict=True)
checks['pdf_two_pages']=len(pdf.pages)==2
checks['pdf_metrics']='903,983' in pdf.pages[0].extract_text() and '105,969' in pdf.pages[0].extract_text()
result={'method':'Independent pandas recomputation of full source; read-only SQLite, openpyxl/XML and strict PDF checks',
        'checks':checks,'all_passed':all(checks.values()),'nativeExcelRecalculationTested':False}
(ROOT/'docs/independent-verification.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
print(json.dumps(result,indent=2))
if not result['all_passed']:raise SystemExit(1)
